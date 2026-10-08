"""Read and write the cost workbook. The workbook is the source of truth for recipes."""
import math
import os
import re
import time
import unicodedata
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

GEN, PRICES, FIXED, CAP = "Costos Generales", "Precios y Márgenes", "Costos Fijos", "Capital Aportado"
GEN_FIRST, GEN_LAST = 6, 35
BLOCK_SLOTS = 16  # ingredient rows in each recipe block


class ExcelBusy(Exception):
    """The workbook is locked, for example because Excel has it open."""


def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


UNITS = {"g": ("m", 1), "kg": ("m", 1000), "ml": ("v", 1), "l": ("v", 1000), "u": ("u", 1), "cda": ("v", 15)}
UNIT_ALIAS = {
    "g": "g", "gr": "g", "gramo": "g", "gramos": "g", "kg": "kg", "kilo": "kg", "kilos": "kg",
    "ml": "ml", "l": "l", "lt": "l", "litro": "l", "litros": "l",
    "u": "u", "un": "u", "unidad": "u", "unidades": "u", "cucharada": "cda", "cucharadas": "cda",
}


def ukey(u):
    return UNIT_ALIAS.get(norm(u))


def conv(qty, a, b):
    """Convert qty from unit a to unit b. Return None when the units do not match."""
    ka, kb = ukey(a), ukey(b)
    if not ka or not kb or UNITS[ka][0] != UNITS[kb][0]:
        return None
    return qty * UNITS[ka][1] / UNITS[kb][1]


def default_insumo_unit(recipe_unit):
    return {"g": "kg", "kg": "kg", "ml": "l", "l": "l", "cda": "l", "u": "unidad"}.get(ukey(recipe_unit), "unidad" if recipe_unit else "kg")


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _qty(ws, r):
    v = ws.cell(r, 2).value
    if isinstance(v, str) and v.startswith("="):
        m = re.fullmatch(r"=\s*([\d.]+)\s*\*\s*B(\d+)", v)
        ref = _num(ws.cell(int(m[2]), 2).value) if m else None
        return (float(m[1]) * ref if ref is not None else None), True
    return _num(v), False


def read_rows(ws, first, last):
    rows = []
    for r in range(first, last + 1):
        qty, formula = _qty(ws, r)
        name, unit = ws.cell(r, 1).value, ws.cell(r, 3).value
        rows.append({
            "row": r, "name": str(name).strip() if name else "", "qty": qty, "qty_formula": formula,
            "unit": str(unit).strip() if unit else "", "cost": _num(ws.cell(r, 4).value),
        })
    return rows


def is_free(row):
    return not row["name"] and row["qty"] is None and not row["qty_formula"]


def line_cost(row):
    return row["qty"] * row["cost"] if row["qty"] is not None and row["cost"] is not None else 0.0


def read_model(wb):
    """Return every value the app needs, computed the same way as the workbook formulas."""
    g = wb[GEN]
    comision, desperdicio = _num(g["B38"].value) or 0.0, _num(g["B39"].value) or 0.0
    gen_rows = read_rows(g, GEN_FIRST, GEN_LAST)
    gen_total = sum(line_cost(r) for r in gen_rows)

    price_cells = {}
    pr = wb[PRICES]
    for r in range(4, 41):
        m = re.fullmatch(r"='([^']+)'!A(\d+)", str(pr.cell(r, 1).value or ""))
        if m:
            price_cells[(m[1], int(m[2]))] = r

    recipes = []
    for sheet in wb.sheetnames:
        if not sheet.startswith("Recetas "):
            continue
        ws, label = wb[sheet], sheet[len("Recetas "):]
        for s in range(1, ws.max_row):
            a = ws.cell(s + 1, 1).value
            if not (isinstance(a, str) and a.startswith("TOTAL INSUMOS ESPEC")):
                continue
            rows = read_rows(ws, s + 3, s + 3 + BLOCK_SLOTS - 1)
            pf, pu = _num(ws.cell(s + 22, 2).value), _num(ws.cell(s + 23, 2).value)
            spec = sum(line_cost(r) for r in rows)
            units = pf * 1000 / pu if pf and pu else 0.0
            unit_cost = (spec + gen_total) / units if units else 0.0
            final = unit_cost / (1 - desperdicio) if desperdicio < 1 else 0.0
            prow = price_cells.get((sheet, s))
            price = _num(pr.cell(prow, 3).value) if prow else None
            margin = pct = food = None
            if price:
                margin = price * (1 - comision) - final
                pct, food = margin / price, final / price
            name = str(ws.cell(s, 1).value or "").strip()
            recipes.append({
                "id": f"{norm(label)}-{s}", "cat": norm(label), "cat_label": label, "sheet": sheet, "start": s,
                "name": name, "rows": rows, "peso_final": pf, "peso_unidad": pu, "units_per_kg": units,
                "spec_total": spec, "lot_total": spec + gen_total, "cost": final, "price": price,
                "price_row": prow, "margin": margin, "margin_pct": pct, "food_cost": food,
                "hidden": norm(name) == "indefinido",
            })

    fx = wb[FIXED]
    fixed = [{"row": r, "name": str(fx.cell(r, 1).value or "").strip(), "amount": _num(fx.cell(r, 2).value) or 0.0,
              "note": str(fx.cell(r, 3).value or "").strip()}
             for r in range(6, 206) if fx.cell(r, 1).value or fx.cell(r, 2).value]
    fixed_total = sum(f["amount"] for f in fixed)

    cp = wb[CAP]
    invest = [{"name": str(cp.cell(r, 1).value), "amount": _num(cp.cell(r, 2).value) or 0.0}
              for r in range(8, 158) if cp.cell(r, 1).value]
    aportes = [{"name": str(cp.cell(r, 1).value), "amount": _num(cp.cell(r, 2).value) or 0.0}
               for r in range(170, 220) if cp.cell(r, 1).value]
    meses = _num(cp["B160"].value) or 0.0
    inv_total = sum(i["amount"] for i in invest)
    needed = inv_total + meses * fixed_total
    contributed = sum(a["amount"] for a in aportes)

    priced = [r for r in recipes if r["price"]]
    avg_margin = sum(r["margin"] for r in priced) / len(priced) if priced else 0.0
    be_month = math.ceil(fixed_total / avg_margin) if avg_margin > 0 else None
    return {
        "params": {"comision": comision, "desperdicio": desperdicio},
        "general": {"rows": gen_rows, "total": gen_total},
        "recipes": recipes,
        "fixed": {"rows": fixed, "total": fixed_total},
        "capital": {"invest": invest, "invest_total": inv_total, "meses": meses, "needed": needed,
                    "aportes": aportes, "contributed": contributed, "diff": contributed - needed},
        "breakeven": {"avg_margin": avg_margin, "month": be_month,
                      "day": math.ceil(be_month / 26) if be_month else None},
    }


def find_block(model, rid):
    if rid == "general":
        return None
    return next((r for r in model["recipes"] if r["id"] == rid), None)


def _efx(r):
    return f'=IF(OR(B{r}="",D{r}=""),0,B{r}*D{r})'


def write_rows(ws, incoming, first, last):
    """Write ingredient rows. A row with no 'row' key goes to the first free slot."""
    current = read_rows(ws, first, last)
    by_row = {c["row"]: c for c in current}
    free = [c["row"] for c in current if is_free(c)]
    keep = set()
    for it in incoming:
        r = it.get("row")
        if r is None:
            if not free:
                raise ValueError("No hay lugar para más ingredientes en esta receta.")
            r = free.pop(0)
        if r not in by_row:
            raise ValueError("Fila de ingrediente no válida.")
        keep.add(r)
        old = by_row[r]
        name = str(it.get("name", "")).strip()
        unit = str(it.get("unit", "")).strip() or old["unit"] or "g"
        qty = it.get("qty")
        qty = float(qty) if qty not in (None, "") else None
        if norm(name) != norm(old["name"]):
            ws.cell(r, 1).value = name
            ws.cell(r, 4).value = None  # the old price belongs to the old ingredient
        if old["qty_formula"] and qty is not None and old["qty"] is not None and abs(qty - old["qty"]) < 1e-9:
            pass
        elif qty != old["qty"]:
            ws.cell(r, 2).value = qty
        if unit != old["unit"]:
            ws.cell(r, 3).value = unit
        ws.cell(r, 5).value = _efx(r)
    for c in current:
        if c["row"] not in keep and c["name"]:
            for col in (1, 2, 4):
                ws.cell(c["row"], col).value = None


def write_recipe(wb, model, rid, rows, peso_final=None, peso_unidad=None, price=None, price_given=False, name=None):
    if rid == "general":
        write_rows(wb[GEN], rows, GEN_FIRST, GEN_LAST)
        return
    rec = find_block(model, rid)
    if not rec:
        raise ValueError("Receta no encontrada.")
    ws, s = wb[rec["sheet"]], rec["start"]
    if name and name.strip():
        ws.cell(s, 1).value = name.strip()
    write_rows(ws, rows, s + 3, s + 3 + BLOCK_SLOTS - 1)
    for off, v in ((22, peso_final), (23, peso_unidad)):
        ws.cell(s + off, 2).value = float(v) if v not in (None, "") else None
    if price_given and rec["price_row"]:
        wb[PRICES].cell(rec["price_row"], 3).value = float(price) if price not in (None, "", 0) else None


BLOCK_LEN = 27  # title row to the final-cost row


def _block_formulas(ws, s):
    g = f"'{GEN}'!"
    ws.cell(s + 1, 5).value = f"=SUM(E{s + 3}:E{s + 18})"
    for r in range(s + 3, s + 19):
        ws.cell(r, 5).value = _efx(r)
    ws.cell(s + 19, 2).value = f"=E{s + 1}"
    ws.cell(s + 20, 2).value = f"={g}E4"
    ws.cell(s + 21, 2).value = f"=B{s + 19}+B{s + 20}"
    ws.cell(s + 24, 2).value = f'=IF(OR(B{s + 22}="",B{s + 23}="",B{s + 23}=0),0,(B{s + 22}*1000)/B{s + 23})'
    ws.cell(s + 25, 2).value = f"=IF(B{s + 24}=0,0,B{s + 21}/B{s + 24})"
    ws.cell(s + 26, 2).value = f"=B{s + 25}/(1-{g}B39)"


def _clear_block(ws, s):
    for r in range(s + 3, s + 19):
        for c in (1, 2, 4):
            ws.cell(r, c).value = None
    ws.cell(s + 22, 2).value = ws.cell(s + 23, 2).value = None


def _add_price_row(wb, sheet, s):
    """Add a row for the recipe in the price sheet. The average row moves down by one."""
    pr = wb[PRICES]
    rows = [r for r in range(4, pr.max_row + 1) if str(pr.cell(r, 1).value or "").startswith("='Recetas")]
    new = rows[-1] + 1
    prom = next(r for r in range(new, pr.max_row + 1) if pr.cell(r, 1).value == "PROMEDIO")
    merged = [str(m) for m in pr.merged_cells.ranges if m.min_row > rows[-1]]
    for m in merged:
        pr.unmerge_cells(m)
    pr.move_range(f"A{new}:F{pr.max_row}", rows=1)
    for m in merged:
        a, b = m.split(":")
        pr.merge_cells(f"{a[0]}{int(a[1:]) + 1}:{b[0]}{int(b[1:]) + 1}")
    prom += 1
    for c in range(1, 7):
        pr.cell(new, c)._style = copy(pr.cell(rows[-1], c)._style)
    pr.cell(new, 1).value = f"='{sheet}'!A{s}"
    pr.cell(new, 2).value = f"='{sheet}'!B{s + 26}"
    com = f"'{GEN}'!B38"
    pr.cell(new, 4).value = f'=IF(OR(C{new}="",C{new}=0),"",C{new}*(1-{com})-B{new})'
    pr.cell(new, 5).value = f'=IF(OR(D{new}="",C{new}="",C{new}=0),"",D{new}/C{new})'
    pr.cell(new, 6).value = f'=IF(OR(C{new}="",C{new}=0),"",B{new}/C{new})'
    for c in "BCDEF":
        pr[f"{c}{prom}"].value = f'=IFERROR(AVERAGEIFS({c}4:{c}{new},$C$4:$C${new},">0"),0)'
    pe = wb["Punto de Equilibrio"]
    for row in pe.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "'Precios y Márgenes'!D" in cell.value:
                cell.value = f"='{PRICES}'!D{prom}"


def create_recipe(wb, model, cat, name):
    """Create a recipe in the sheet of the meat cat. Reuse an empty block when one exists."""
    sheet = next((r["sheet"] for r in model["recipes"] if r["cat"] == norm(cat)), None) or f"Recetas {cat}"
    if sheet not in wb.sheetnames:
        raise ValueError("No existe esa carne en el Excel.")
    ws = wb[sheet]
    blocks = [r for r in model["recipes"] if r["sheet"] == sheet]
    for b in blocks:
        empty = not any(x["name"] for x in b["rows"]) and not b["peso_final"] and not b["peso_unidad"]
        if empty and norm(b["name"]) in ("", "indefinido"):
            ws.cell(b["start"], 1).value = name
            return f"{norm(cat)}-{b['start']}"
    tpl = max(b["start"] for b in blocks)
    s = tpl + 29
    for off in range(BLOCK_LEN):
        if ws.row_dimensions[tpl + off].height:
            ws.row_dimensions[s + off].height = ws.row_dimensions[tpl + off].height
        for c in range(1, 7):
            src, dst = ws.cell(tpl + off, c), ws.cell(s + off, c)
            dst._style = copy(src._style)
            if c == 1 and off in (1, 2, 19, 20, 21, 22, 23, 24, 25, 26):
                dst.value = src.value
            elif off == 2 and c in (2, 3, 4, 5):
                dst.value = src.value
    ws.merge_cells(start_row=s, start_column=1, end_row=s, end_column=6)
    ws.merge_cells(start_row=s + 1, start_column=1, end_row=s + 1, end_column=4)
    for r in range(s + 3, s + 19):
        ws.cell(r, 3).value = "g"
    ws.cell(s, 1).value = name
    _block_formulas(ws, s)
    _add_price_row(wb, sheet, s)
    return f"{norm(cat)}-{s}"


def delete_recipe(wb, model, rid):
    rec = find_block(model, rid)
    ws, s = wb[rec["sheet"]], rec["start"]
    _clear_block(ws, s)
    ws.cell(s, 1).value = "Indefinido"
    if rec["price_row"]:
        wb[PRICES].cell(rec["price_row"], 3).value = None


def write_params(wb, comision, desperdicio):
    wb[GEN]["B38"].value, wb[GEN]["B39"].value = comision, desperdicio


def write_fixed(wb, rows):
    ws = wb[FIXED]
    for r in range(6, 206):
        for c in (1, 2, 3):
            ws.cell(r, c).value = None
    for i, it in enumerate(rows[:200]):
        r = 6 + i
        ws.cell(r, 1).value = str(it.get("name", "")).strip()
        ws.cell(r, 2).value = float(it.get("amount") or 0)
        ws.cell(r, 3).value = str(it.get("note", "")).strip() or None


def sync_prices(wb, model, resolve):
    """Copy the price of each linked insumo into the cost column. Return True if a cell changed.

    resolve(name) returns (insumo_unit, insumo_price) or None.
    """
    changed = False
    blocks = [(wb[GEN], model["general"]["rows"])] + [(wb[r["sheet"]], r["rows"]) for r in model["recipes"]]
    for ws, rows in blocks:
        for row in rows:
            if not row["name"]:
                continue
            ins = resolve(row["name"])
            if not ins or not ins[1] or not row["unit"]:
                continue
            per_unit = conv(1, ins[0], row["unit"])  # recipe units in one insumo unit
            if not per_unit:
                continue
            cost = round(ins[1] / per_unit, 6)
            r = row["row"]
            if row["cost"] is None or abs(row["cost"] - cost) > 1e-9:
                ws.cell(r, 4).value = cost
                changed = True
            if ws.cell(r, 5).value is None:
                ws.cell(r, 5).value = _efx(r)
                changed = True
    return changed


class ExcelBook:
    def __init__(self, path, backup_dir):
        self.path, self.backup_dir = Path(path), Path(backup_dir)

    def open(self):
        return load_workbook(self.path)

    def _backup(self):
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        files = sorted(self.backup_dir.glob("Costos-*.xlsx"))
        if files and time.time() - files[-1].stat().st_mtime < 600:
            return
        (self.backup_dir / time.strftime("Costos-%Y%m%d-%H%M%S.xlsx")).write_bytes(self.path.read_bytes())
        for old in sorted(self.backup_dir.glob("Costos-*.xlsx"))[:-40]:
            old.unlink()

    def save(self, wb):
        self._backup()
        wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula on open
        tmp = self.path.with_suffix(".tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError as e:
            raise ExcelBusy("No se pudo guardar el Excel. Cerrá el archivo si lo tenés abierto en Excel.") from e
